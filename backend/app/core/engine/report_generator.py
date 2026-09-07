import os
import csv
from typing import List, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from backend.app.schemas.types import DeviationItem, AnalysisSummary, SeverityLevel
from backend.app.config import SNIPPETS_DIR

class AuditReportGenerator:
    """
    Industrial Deviation Audit Report Generator.
    Produces:
    1. Professional PDF Audit Report (ReportLab)
    2. Formatted Excel (XLSX) Workbook
    3. Standard CSV Export
    """

    def generate_pdf_report(
        self,
        summary: AnalysisSummary,
        deviations: List[DeviationItem],
        output_pdf_path: str
    ):
        doc = SimpleDocTemplate(
            output_pdf_path,
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Custom styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#0F172A')
        )
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            leading=15,
            textColor=colors.HexColor('#475569')
        )
        h2_style = ParagraphStyle(
            'SectionH2',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#1E293B'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#334155')
        )
        cell_bold = ParagraphStyle(
            'CellBold',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#0F172A')
        )
        cell_regular = ParagraphStyle(
            'CellRegular',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor('#334155')
        )

        story = []

        # --- Header ---
        story.append(Paragraph("DEVIATION INTELLIGENCE", title_style))
        story.append(Paragraph("Automated Engineering Control-System Deviation Audit Report", subtitle_style))
        story.append(Spacer(1, 10))

        # --- Project Meta Table ---
        meta_data = [
            [
                Paragraph("<b>Project:</b> " + (summary.project_name or "Industrial Control System"), body_style),
                Paragraph("<b>Analysis ID:</b> " + summary.id[:8], body_style),
                Paragraph("<b>Generated:</b> " + datetime.now().strftime("%Y-%m-%d %H:%M"), body_style)
            ],
            [
                Paragraph("<b>PDF Drawing:</b> " + summary.pdf_filename, body_style),
                Paragraph("<b>PDF Rev:</b> " + (summary.pdf_revision or "Rev-0"), body_style),
                Paragraph("<b>Pages:</b> " + str(summary.stats.page_count), body_style)
            ],
            [
                Paragraph("<b>Config Export:</b> " + summary.config_filename, body_style),
                Paragraph("<b>Config Rev:</b> " + (summary.config_revision or "Rev-0"), body_style),
                Paragraph("<b>Parser Profile:</b> " + (summary.parser_profile or "Generic"), body_style)
            ]
        ]
        meta_table = Table(meta_data, colWidths=[200, 160, 160])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 14))

        # --- Executive Summary Statistics ---
        story.append(Paragraph("Executive Summary & Risk Metrics", h2_style))
        stats = summary.stats
        stats_data = [
            ["TOTAL DEVIATIONS", "CRITICAL", "HIGH", "MEDIUM", "LOW", "REVIEW REQUIRED"],
            [
                str(stats.total_deviations),
                str(stats.critical_count),
                str(stats.high_count),
                str(stats.medium_count),
                str(stats.low_count),
                str(stats.review_required_count)
            ]
        ]
        stats_table = Table(stats_data, colWidths=[90, 85, 85, 85, 85, 90])
        stats_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F172A')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 8),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
            ('TOPPADDING', (0,0), (-1,0), 6),
            ('BACKGROUND', (0,1), (0,1), colors.HexColor('#F1F5F9')),
            ('BACKGROUND', (1,1), (1,1), colors.HexColor('#FEE2E2')), # Red
            ('BACKGROUND', (2,1), (2,1), colors.HexColor('#FFEDD5')), # Orange
            ('BACKGROUND', (3,1), (3,1), colors.HexColor('#FEF9C3')), # Yellow
            ('BACKGROUND', (4,1), (4,1), colors.HexColor('#E0F2FE')), # Blue
            ('BACKGROUND', (5,1), (5,1), colors.HexColor('#F3E8FF')), # Purple
            ('FONTNAME', (0,1), (-1,1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,1), (-1,1), 12),
            ('TOPPADDING', (0,1), (-1,1), 8),
            ('BOTTOMPADDING', (0,1), (-1,1), 8),
            ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ]))
        story.append(stats_table)
        story.append(Spacer(1, 14))

        # --- Detailed Deviation Findings ---
        story.append(Paragraph("Detailed Deviation Findings", h2_style))

        # Table header
        table_rows = [[
            Paragraph("<b>ID / Page</b>", cell_bold),
            Paragraph("<b>Type & Severity</b>", cell_bold),
            Paragraph("<b>PDF Drawing vs Configuration</b>", cell_bold),
            Paragraph("<b>Engineering Explanation & Action</b>", cell_bold),
            Paragraph("<b>Conf</b>", cell_bold)
        ]]

        for dev in deviations:
            sev_color = "#DC2626" if dev.severity == SeverityLevel.CRITICAL else (
                "#EA580C" if dev.severity == SeverityLevel.HIGH else (
                    "#CA8A04" if dev.severity == SeverityLevel.MEDIUM else "#2563EB"
                )
            )

            id_p = Paragraph(f"<b>{dev.deviation_number}</b><br/>Page {dev.page_number}<br/><font color='#64748B'>{dev.equipment_tag}</font>", cell_regular)
            type_p = Paragraph(f"<font color='{sev_color}'><b>{dev.severity.value}</b></font><br/>{dev.type.value.replace('_', ' ')}", cell_regular)
            evidence_p = Paragraph(f"<b>PDF:</b> {dev.pdf_expected}<br/><b>CFG:</b> {dev.config_actual}<br/><font color='#64748B'>Lines {dev.config_evidence.line_start}-{dev.config_evidence.line_end}</font>", cell_regular)
            explain_p = Paragraph(f"{dev.explanation}<br/><b>Action:</b> <font color='#0284C7'>{dev.recommended_action}</font>", cell_regular)
            conf_p = Paragraph(f"<b>{dev.confidence}%</b>", cell_regular)

            table_rows.append([id_p, type_p, evidence_p, explain_p, conf_p])

        dev_table = Table(table_rows, colWidths=[65, 95, 145, 180, 35])
        dev_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING', (0,0), (-1,-1), 5),
            ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(dev_table)

        doc.build(story)

    def generate_excel_report(
        self,
        summary: AnalysisSummary,
        deviations: List[DeviationItem],
        output_xlsx_path: str
    ):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Deviations"

        # Headers
        headers = [
            "Deviation ID", "Page", "Severity", "Deviation Type", "Status",
            "Equipment / Block", "PDF Expected Value", "Configuration Value",
            "Confidence (%)", "Explanation", "Recommended Action",
            "Config Line Start", "Config Line End", "Review Notes"
        ]

        # Header styling
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Fill rows
        for row_idx, dev in enumerate(deviations, start=2):
            row_data = [
                dev.deviation_number,
                dev.page_number,
                dev.severity.value,
                dev.type.value,
                dev.status.value,
                dev.block_name or dev.equipment_tag,
                dev.pdf_expected,
                dev.config_actual,
                dev.confidence,
                dev.explanation,
                dev.recommended_action,
                dev.config_evidence.line_start,
                dev.config_evidence.line_end,
                dev.reviewer_notes or ""
            ]
            ws.append(row_data)

            # Apply borders and cell formatting
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)

        # Set column widths
        col_widths = [14, 8, 12, 22, 12, 18, 25, 25, 12, 45, 45, 14, 14, 25]
        for idx, width in enumerate(col_widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = width

        wb.save(output_xlsx_path)

    def generate_csv_report(
        self,
        deviations: List[DeviationItem],
        output_csv_path: str
    ):
        with open(output_csv_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow([
                "Deviation ID", "Page", "Severity", "Deviation Type", "Status",
                "Block", "PDF Expected", "Config Actual", "Confidence",
                "Explanation", "Recommended Action", "Line Start", "Line End"
            ])
            for dev in deviations:
                writer.writerow([
                    dev.deviation_number,
                    dev.page_number,
                    dev.severity.value,
                    dev.type.value,
                    dev.status.value,
                    dev.block_name or dev.equipment_tag,
                    dev.pdf_expected,
                    dev.config_actual,
                    dev.confidence,
                    dev.explanation,
                    dev.recommended_action,
                    dev.config_evidence.line_start,
                    dev.config_evidence.line_end
                ])
